"""Writes the training lesson decks (training/decks/*.pptx) from the content below.

    uv run --with python-pptx==1.0.2 python training/decks/build_decks.py

The decks are real PowerPoint files: the app imports them with the same deck importer
as a church's own, so they double as fidelity fixtures. Each slide's speaker notes hold
the trainer's script, and in Training Session 2 the surprises the trainer calls, so a trainer
reads them on the confidence monitor or beside the item on air in Present.

Builds are ordinary PowerPoint "appear by paragraph" click effects, written as the XML
PowerPoint writes, and some slides have a Fade transition. Text uses Noto Sans, which
Lampstand bundles, so a deck looks the same on every computer.

Content is plain words for volunteers aged 15 to 70. Scripture quoted is the King James
Version (public domain). Nothing here is under licence.
"""

import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parent
FONT = "Noto Sans"
NAVY = RGBColor(0x0B, 0x0E, 0x16)
INK = RGBColor(0xF5, 0xF1, 0xE8)
GOLD = RGBColor(0xE9, 0xB9, 0x49)
MUTED = RGBColor(0xA8, 0xA2, 0x96)
KEYCAP = RGBColor(0x22, 0x27, 0x36)
W, H = Emu(12192000), Emu(6858000)  # 16:9, 13.333 x 7.5 in
M = Emu(685800)  # 0.75 in margin


def text(slide, x, y, w, h, runs, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """A text box. `runs` is a string or a list of paragraphs."""
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, line in enumerate(paras):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.5)
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def footer(slide, label):
    text(slide, M, H - Emu(500000), W - 2 * M, Emu(300000), label, 12, MUTED)


def background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def builds(slide, shape, count):
    """One click per paragraph of `shape`: PowerPoint's Appear effect, by paragraph."""
    sid = shape.shape_id
    clicks = []
    n = 3
    for p in range(count):
        a, b, c, d = n, n + 1, n + 2, n + 3
        n += 4
        clicks.append(
            f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{c}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{sid}"><p:txEl><p:pRg st="{p}" end="{p}"/></p:txEl></p:spTgt></p:tgtEl>'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="visible"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        )
    xml = (
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        + "".join(clicks)
        + '</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst><p:bldP spid="{sid}" grpId="0" build="p"/></p:bldLst></p:timing>'
    )
    slide._element.append(etree.fromstring(xml))


def fade(slide):
    xml = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>'
    timing = slide._element.find(qn("p:timing"))
    el = etree.fromstring(xml)
    if timing is not None:
        timing.addprevious(el)
    else:
        slide._element.append(el)


def keycap(slide, x, y, key, w=Emu(2400000)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Emu(560000))
    box.fill.solid()
    box.fill.fore_color.rgb = KEYCAP
    box.line.color.rgb = GOLD
    box.line.width = Pt(1.5)
    frame = box.text_frame
    frame.margin_left = frame.margin_right = Emu(60000)
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = key
    r.font.name = FONT
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = GOLD


def add(prs, spec, label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    background(slide)
    kind = spec["kind"]
    if kind == "title":
        if spec.get("kicker"):
            text(slide, M, Emu(1700000), W - 2 * M, Emu(500000), spec["kicker"], 20, GOLD, bold=True)
        text(slide, M, Emu(2250000), W - 2 * M, Emu(1500000), spec["heading"], 48, INK, bold=True)
        if spec.get("sub"):
            text(slide, M, Emu(3900000), W - 2 * M, Emu(1400000), spec["sub"], 24, MUTED)
    elif kind == "bullets":
        text(slide, M, Emu(600000), W - 2 * M, Emu(900000), spec["heading"], 36, GOLD, bold=True)
        body = text(slide, M, Emu(1750000), W - 2 * M, Emu(4300000), ["•  " + i for i in spec["items"]], 28)
        if spec.get("builds"):
            builds(slide, body, len(spec["items"]))
    elif kind == "keys":
        text(slide, M, Emu(600000), W - 2 * M, Emu(900000), spec["heading"], 36, GOLD, bold=True)
        y = Emu(1700000)
        step = Emu(min(760000, int(4300000 / max(1, len(spec["rows"])))))
        for key, what in spec["rows"]:
            keycap(slide, M, y, key)
            text(slide, M + Emu(2700000), y + Emu(60000), W - 2 * M - Emu(2700000), Emu(560000), what, 24, INK)
            y += step
    elif kind == "quote":
        text(slide, M + Emu(400000), Emu(1600000), W - 2 * M - Emu(800000), Emu(2800000), spec["text"], 36, INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, M, Emu(4600000), W - 2 * M, Emu(500000), spec["source"], 22, GOLD, align=PP_ALIGN.CENTER)
    else:
        raise ValueError(kind)
    if spec.get("fade"):
        fade(slide)
    footer(slide, label)
    slide.notes_slide.notes_text_frame.text = spec["notes"].strip()
    return slide


def save(name, title, label, slides):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for spec in slides:
        add(prs, spec, label)
    props = prs.core_properties
    props.title = title
    props.author = "Lampstand"
    props.created = props.modified = datetime(2026, 10, 1)
    props.last_modified_by = "Lampstand"
    props.revision = 1
    raw = io.BytesIO()
    prs.save(raw)
    # Same content, same bytes: fixed zip dates, so a rebuild changes nothing.
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, src.read(info.filename))
    (OUT / name).write_bytes(out.getvalue())
    print(f"{name}: {len(slides)} slides, {len(out.getvalue()) // 1024} KB")


# ---------------------------------------------------------------- Training Session 0

T0 = "Lampstand training · 0 Orientation"
save("t0-plan-and-present.pptx", "Plan and Present", T0, [
    {"kind": "title", "kicker": "Training Session 0", "heading": "Two modes and a toolbox",
     "sub": "Plan during the week. Present on Sunday. Setup once.",
     "notes": """
Trainer: Welcome everyone. This takes five minutes and is for everybody on the team,
not only the people at the projection computer.
You are watching a slide deck that Lampstand is showing, just as it shows a sermon deck.
Press Space to move on."""},
    {"kind": "bullets", "heading": "Plan: during the week", "builds": True, "fade": True, "items": [
        "Put the service in order",
        "Add songs, Bible readings, slides and videos",
        "Anyone on the team can help, from their own computer",
    ], "notes": """
Trainer: Plan is where the service gets built. Worship leaders, pastors and admins
mostly live here.
Each press of Space brings in the next point. That is called a build.
Say: nothing you do in Plan reaches the screen in the room."""},
    {"kind": "bullets", "heading": "Present: on Sunday", "builds": True, "fade": True, "items": [
        "Open one service for live",
        "Move through it with Space",
        "The room sees only what you send",
    ], "notes": """
Trainer: Present is the operator's screen. One service is open at a time, and every slide
of it is laid out as cards. The card marked LIVE, in red, is what the room sees, and so is the
small picture at the top left. Everything else is only for you.
This training service is open in Present right now, and it works like any other service."""},
    {"kind": "bullets", "heading": "Setup: once, then rarely", "fade": True, "items": [
        "Screens: which screen is the projector",
        "Themes: how the words look",
        "Devices and Team: who does what, on which computer",
    ], "notes": """
Trainer: Setup is for the media or tech volunteer. It is set once and then left alone.
If something looks wrong on the projector on Sunday, it is usually a Setup question,
and it can wait until after the service. Blank the screen and carry on."""},
    {"kind": "keys", "heading": "Move between them", "fade": True, "rows": [
        ("Ctrl+1", "Plan  (⌘1 on a Mac)"),
        ("Ctrl+2", "Present  (⌘2 on a Mac)"),
        ("Ctrl+\\", "Switch between Plan and Present"),
    ], "notes": """
Trainer: Try it now. Press Ctrl and backslash (Command and backslash on a Mac) to go to
Plan. Find this service in the list under Training. Then press it again to come back.
The service stays open for live while you look at Plan."""},
])

# ---------------------------------------------------------------- Training Session 1

T1 = "Lampstand training · 1 Run a service"
save("t1-present-screen.pptx", "The Present screen", T1, [
    {"kind": "title", "kicker": "Training Session 1", "heading": "The Present screen",
     "sub": "Four places to look. One rule.",
     "notes": """
Trainer: Ten minutes from now you will run a short practice service with no wrong slide
and no desktop showing. Point at each part of the screen as we go."""},
    {"kind": "bullets", "heading": "Four places to look", "builds": True, "items": [
        "The wall: every slide of the service",
        "Red LIVE: what the room sees now",
        "Gold NEXT: the next slide",
        "The strip on top: the main screen and the buttons",
    ], "notes": """
Trainer: Point at each one on the real screen as it appears.
The wall is every slide, item by item, in order. The red LIVE card is on air, and the gold
NEXT card is the next slide. On a slide with builds, Space plays the next build first.
The notes, where you are reading this, sit beside the item on air. The strip on top has a
small picture of the main screen, then Focus, Quick add, Previous, Next and Blank. Show, Idle and Blank, top right, switch the room between the slide, the idle
screen and black. Press F for Focus to see only the cards; F or Esc brings the rest back."""},
    {"kind": "quote", "text": "Nothing reaches the screen until you send it.", "source": "The one rule",
     "fade": True, "notes": """
Trainer: Choosing an item with Shift and Down only marks it for you. The room sees nothing new. It goes on air when you
start it with Enter, or click one of its cards. Space always moves on from what is live."""},
    {"kind": "keys", "heading": "The keys you need today", "rows": [
        ("Space  →", "Next slide"),
        ("←", "Previous slide"),
        (".", "Blank the screen, and back"),
        ("⇧↓  Enter", "Choose the next item, then start it"),
    ], "notes": """
Trainer: These four are enough to run a whole service. A presenter clicker sends the same
keys. The next video shows them in action."""},
])

save("t1-sermon-practice.pptx", "Practice sermon: The Lord is my shepherd", T1, [
    {"kind": "title", "kicker": "Psalm 23", "heading": "The Lord is my shepherd",
     "sub": "A practice sermon in five slides",
     "notes": """
Preacher: Good morning. Our reading today is Psalm 23.
Operator: this deck has builds. Each Space press brings in one point. When the slide is
complete, the next Space goes to the next slide. While a slide with builds is on air, a panel
above the wall shows what is on screen beside the next click, and how many clicks are left."""},
    {"kind": "bullets", "heading": "He leads me", "builds": True, "fade": True, "items": [
        "Green pastures",
        "Still waters",
        "Paths of righteousness",
    ], "notes": """
Preacher: He makes me lie down in green pastures. He leads me beside the still waters.
He leads me in the paths of righteousness.
Operator: three points, three presses. If you press one too many, press Left to step back."""},
    {"kind": "quote", "text": "Yea, though I walk through the valley of the shadow of death, I will fear no evil: for thou art with me.",
     "source": "Psalm 23:4 (KJV)", "fade": True, "notes": """
Preacher: Read the verse slowly.
Operator: a quote slide has no builds. Leave it up while the preacher talks about it."""},
    {"kind": "bullets", "heading": "Three things to take home", "builds": True, "items": [
        "You are known",
        "You are led",
        "You are not alone",
    ], "notes": """
Preacher: Three things to take home this week.
Operator: if the preacher jumps ahead ("skip to the last slide"), click that slide's card on
the wall. It goes on air straight away."""},
    {"kind": "quote", "text": "Surely goodness and mercy shall follow me all the days of my life.",
     "source": "Psalm 23:6 (KJV)", "fade": True, "notes": """
Preacher: Let us pray.
Operator: this is the last slide. When the prayer ends, press Space to move to the next item,
or Idle if the room should see the idle screen."""},
])

save("t1-ready.pptx", "You're ready", T1, [
    {"kind": "title", "kicker": "Training Session 1", "heading": "You're ready",
     "sub": "You have run a service: songs, scripture, a sermon deck and the blank screen.",
     "notes": """
Trainer: Well done. Ask the trainee: what would you do if the wrong slide went up?
(Answer: press full stop to blank, find the right slide, then bring it back.)"""},
    {"kind": "bullets", "heading": "Sunday, step by step", "builds": True, "items": [
        "Before: screens on, open the service for live",
        "Start on the idle screen; Space puts the first slide up",
        "During: Space to move on, full stop to blank",
        "After: Close service, top right",
    ], "notes": """
Trainer: Walk through a real Sunday. Point out Close service in the bar at the top right.
Give them the printed key sheet to keep by the computer.
Next: Training Session 2, Confident operator, practises the five surprises that happen on a real Sunday."""},
])

# ---------------------------------------------------------------- Training Session 2

T2 = "Lampstand training · 2 Confident operator"
save("t2-briefing.pptx", "Briefing: five surprises", T2, [
    {"kind": "title", "kicker": "Training Session 2", "heading": "Five surprises",
     "sub": "A rehearsal service. The trainer calls the surprises. You keep the screen calm.",
     "notes": """
Trainer: This takes about thirty minutes. The surprises are written in the notes of each
item in the rehearsal service below, so read them beside each item as it comes on air.
Do not tell the trainee which surprise is coming. Run the rehearsal twice: once slowly,
then again at the speed of a real service. Press Reset training in Plan between runs."""},
    {"kind": "keys", "heading": "The five moves", "rows": [
        ("1  …  9   C", "The leader repeats a verse or the chorus"),
        ("⇧↓  Enter", "The order changes"),
        ("/", "Scripture nobody planned"),
        ("⇧↓  Enter", "An updated deck arrives"),
        (".", "Emergency: blank first, think second"),
    ], "notes": """
Trainer: Go through each move once before the rehearsal.
1 to 9 jump to that verse of the live song; C jumps to the chorus.
Shift and Down choose another item; Enter starts it.
Slash opens Quick add: type a reference like Rom 8:28 and press Shift and Enter to show it now.
Full stop blanks the screen. Press it again to bring the slide back."""},
    {"kind": "quote", "text": "The room never sees you looking for something.", "source": "The aim",
     "notes": """
Trainer: The skill is to blank or hold while you look, never to search on the live screen.
Choosing with Shift and the arrows only marks an item for you. The room sees nothing new until you press Enter."""},
])

SERMON_V1 = [
    {"kind": "title", "kicker": "Rehearsal sermon", "heading": "The Lord is my shepherd",
     "sub": "As sent on Thursday",
     "notes": """
SURPRISE 4 (updated deck). The trainer gave the cue during the reading:
"The preacher changed the verse slide this morning. Use the new deck."
Operator: if this deck is on air, blank first. Do not use this one. Shift and Down to the next item, "Sermon (updated Sunday
morning)", and press Enter. On a real Sunday the planner adds the new file in Plan, from any
computer on the team, and it appears in your order. Never edit during the service."""},
    {"kind": "bullets", "heading": "He leads me", "builds": True, "items": [
        "Green pastures", "Still waters", "Paths of righteousness"],
     "notes": "Preacher: He leads me. Three points, three presses."},
    {"kind": "quote", "text": "Yea, though I walk through the valley of the shadow of death, I will fear no evil.",
     "source": "Psalm 32:4", "notes": "Old version: this reference is wrong (it should be Psalm 23:4)."},
]
save("t2-sermon-thursday.pptx", "Sermon (sent on Thursday)", T2, SERMON_V1)

SERMON_V2 = [
    {**SERMON_V1[0], "sub": "Updated Sunday morning", "notes": """
Preacher: Good morning. This is the corrected deck.
Operator: well done for choosing it. Space through the builds as usual."""},
    SERMON_V1[1],
    {"kind": "quote", "text": "Yea, though I walk through the valley of the shadow of death, I will fear no evil: for thou art with me.",
     "source": "Psalm 23:4 (KJV)", "fade": True, "notes": """
Preacher: Read the verse. This is the slide that changed.
SURPRISE 5 (blank in an emergency) can come any time in this deck. Trainer: say "Blank!"
Operator: full stop. Then full stop again when the trainer says "Back"."""},
    {"kind": "quote", "text": "Surely goodness and mercy shall follow me all the days of my life.",
     "source": "Psalm 23:6 (KJV)", "notes": "Preacher: Let us pray. Operator: Space to move on when the prayer ends."},
]
save("t2-sermon-sunday.pptx", "Sermon (updated Sunday morning)", T2, SERMON_V2)

save("t2-phones.pptx", "Phones and tablets", T2, [
    {"kind": "title", "kicker": "Training Session 2", "heading": "Phones and tablets",
     "sub": "A helper can run the service from a phone. A speaker can step through their own slides.",
     "notes": """
Trainer: This part takes about ten minutes and needs a phone on the same Wi-Fi as this computer.
Phones and tablets connect from Present, with nothing to install. There are two kinds, and
choosing the right one is the whole lesson."""},
    {"kind": "bullets", "heading": "Two kinds of phone", "builds": True, "items": [
        "Run the service: every slide as cards, Next and Back, Blank, Idle and the sound",
        "Speaker: one item's slides, with its notes and a timer",
        "A speaker's Next and Back work only while their slides are on the screen",
        "You can still run everything from this computer",
    ], "notes": """
Trainer: Run the service is for a second volunteer, or for you away from the desk. It can do
almost everything Present can, so give it only to someone on the team.
Speaker is for the preacher or anyone presenting their own slides. Their phone shows only their
item: the slide on screen, the next click, their notes and a timer. Next on their last slide does
nothing, so a speaker can never move the service on by accident. If you clear the screen, their
phone says so and waits."""},
    {"kind": "bullets", "heading": "Connecting one", "builds": True, "items": [
        "Present › More › Connect a phone or tablet",
        "Choose Run the service, or Speaker and whose slides",
        "They scan the code and type the PIN",
        "Setup › Devices lists every phone; Revoke disconnects one",
    ], "notes": """
Trainer: The PIN works for five minutes and only once. A speaker's phone stops working by itself
when the service is closed. A Run the service phone stays paired until you revoke it in Setup ›
Devices. Now try it: the next item is the practice."""},
])

save("t2-debrief.pptx", "Debrief", T2, [
    {"kind": "bullets", "heading": "How did it go?", "builds": True, "items": [
        "Did the room ever see a wrong slide?",
        "Did the room ever see the desktop?",
        "Which surprise felt hardest?",
    ], "notes": """
Trainer: Ask each question and let the trainee answer. If any answer is yes, run the
rehearsal again: press Reset training in Plan, open Training Session 2 for live and start from the top.
A confident operator gets through all five surprises with the screen calm every time."""},
    {"kind": "title", "kicker": "Training Session 2", "heading": "Confident operator",
     "sub": "Next for planners: Training Session 3. For media volunteers: Training Session 4.",
     "notes": """
Trainer: That is the end of operator training. Training Session 3 (Planner) and Training Session 4 (Setup and the
stream) come in a later version of Lampstand. Close the service with Close service, top right."""},
])

if __name__ == "__main__":
    sys.exit(0)
